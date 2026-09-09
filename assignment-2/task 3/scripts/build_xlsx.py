"""
Build task3_passing.xlsx — Excel workbook version of Task 3 (Passing/Distribution).

Sheets:
  README            - question, method, data source, deviations
  Data (n=450)      - full wrangled population, Cmp computed by formula
  Sample (n=60)     - the exact random sample used (30 Finalist + 30 Non-finalist,
                       Finalist rows first so descriptive stats can use plain ranges)
  Descriptive Stats - AVERAGE/MEDIAN/STDEV/MIN/MAX formulas per group
  Inferential Stats - 95% CI (Welch-Satterthwaite) and Welch's t-test, all by formula
  Answer Chart       - the headline PNG embedded for reference

All statistics are live formulas over the Sample sheet, not pasted results.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

RANDOM_SEED = 398629
FONT_NAME = "Arial"

HEADER_FILL = PatternFill("solid", fgColor="2A78D6")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10.5)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="52514E")
BODY_FONT = Font(name=FONT_NAME, size=10.5)
BOLD_FONT = Font(name=FONT_NAME, bold=True, size=10.5)
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9.5, color="898781")
THIN = Side(style="thin", color="E1E0D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------
# Load data and reproduce the exact sample used in task3_passing.py
# ---------------------------------------------------------------------
df = pd.read_csv("data/wc2026_passing.csv")
df = df.dropna(subset=["Player", "Squad", "Att", "Cmp_pct", "Group"]).drop_duplicates(
    subset=["Player", "Squad"]
)
df = df[df["Att"] >= 20].copy()  # NOTE: no re-sort here — must match task3_passing.py's
                                  # row order exactly, since pandas .sample(random_state=)
                                  # depends on row order, not just the seed.

fin_sample = (
    df[df.Group == "Finalist"]
    .sample(n=30, random_state=RANDOM_SEED)
    .sort_values(["Squad", "Player"])  # sort AFTER sampling — display order only
)
non_sample = (
    df[df.Group == "Non-finalist"]
    .sample(n=30, random_state=RANDOM_SEED)
    .sort_values(["Squad", "Player"])  # sort AFTER sampling — display order only
)
sample = pd.concat([fin_sample, non_sample]).reset_index(drop=True)

# Population sheet: sort for readability only, after the sample above was already drawn.
df = df.sort_values(["Group", "Squad", "Player"]).reset_index(drop=True)

wb = Workbook()

# ---------------------------------------------------------------------
# Sheet: README
# ---------------------------------------------------------------------
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
autosize(ws, [110])

r = 1
ws.cell(row=r, column=1, value="Task 3 — Passing / Distribution").font = TITLE_FONT
r += 1
ws.cell(row=r, column=1, value="Author: Sudip Sunar (S398629)").font = SUBTITLE_FONT
r += 2

lines = [
    ("Analytic question", BOLD_FONT),
    ("Did players from the two 2026 FIFA World Cup finalist teams (Spain and "
     "Argentina — Spain won the final 1-0 after extra time) achieve a "
     "significantly higher pass completion rate (Cmp%) than players from teams "
     "that did not reach the final?", BODY_FONT),
    ("", BODY_FONT),
    ("Data source", BOLD_FONT),
    ("FIFA's official Player Statistics page for the 2026 FIFA World Cup "
     "(Distribution → Passes): fifa.com/en/tournaments/mens/worldcup/"
     "canadamexicousa2026/statistics/player-statistics?group=gcp_distribution&stat=passes", BODY_FONT),
    ("Originally planned to use FBref's Passing table; FBref sits behind a "
     "Cloudflare bot-verification CAPTCHA that was not solved programmatically, "
     "so the real per-player figures were pulled from FIFA's own official stats "
     "page instead — same tournament, same core statistic (passes attempted + "
     "accuracy %).", NOTE_FONT),
    ("", BODY_FONT),
    ("Method", BOLD_FONT),
    ("1. Data wrangling — drop missing values and duplicate player/squad rows; "
     "drop players with fewer than 20 passes attempted (stand-in for a minutes-"
     "played filter, since FIFA's table does not expose minutes/90s).", BODY_FONT),
    ("2. Sampling — split into Finalist (ESP, ARG) vs Non-finalist; simple random "
     "sample of up to 30 players per group (seed = 398629, student ID).", BODY_FONT),
    ("3. Descriptive statistics — mean, median, std dev, min, max Cmp% per group.", BODY_FONT),
    ("4. 95% CI for the difference in mean Cmp% (Welch-Satterthwaite).", BODY_FONT),
    ("5. Welch's two-sample t-test (unequal variances), alpha = 0.05.", BODY_FONT),
    ("", BODY_FONT),
    ("How this workbook is built", BOLD_FONT),
    ("'Data (n=450)' holds the full wrangled population. 'Sample (n=60)' holds "
     "the exact 30+30 random draw analysed (Finalist rows first, then Non-"
     "finalist, so every downstream formula can reference a plain range). "
     "'Descriptive Stats' and 'Inferential Stats' compute everything live with "
     "formulas over the Sample sheet — change a Cmp_pct value there and every "
     "number recalculates.", BODY_FONT),
]
for text, font in lines:
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 14 if not text else max(14, 14 * (len(text) // 100 + 1))
    r += 1

# ---------------------------------------------------------------------
# Sheet: Data (n=450)
# ---------------------------------------------------------------------
ws = wb.create_sheet("Data (n=450)")
headers = ["Player", "Squad", "Pos", "Att", "Cmp_pct", "Cmp (derived)", "Group"]
for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
style_header_row(ws, 1, len(headers))
ws.freeze_panes = "A2"

for i, row in df.iterrows():
    r = i + 2
    ws.cell(row=r, column=1, value=row["Player"]).font = BODY_FONT
    ws.cell(row=r, column=2, value=row["Squad"]).font = BODY_FONT
    ws.cell(row=r, column=3, value=row["Pos"]).font = BODY_FONT
    ws.cell(row=r, column=4, value=int(row["Att"])).font = BODY_FONT
    ws.cell(row=r, column=5, value=int(row["Cmp_pct"])).font = BODY_FONT
    ws.cell(row=r, column=6, value=f"=ROUND(D{r}*E{r}/100,0)").font = BODY_FONT
    ws.cell(row=r, column=7, value=f'=IF(OR(B{r}="ESP",B{r}="ARG"),"Finalist","Non-finalist")').font = BODY_FONT
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = BORDER

autosize(ws, [22, 8, 6, 8, 10, 14, 14])
last = len(df) + 1
note_row = last + 2
ws.cell(row=note_row, column=1,
        value=f"Source: FIFA official 2026 World Cup Player Statistics (Passes/Distribution). "
              f"n={len(df)} players with Att >= 20 (all players pulled already clear this bar). "
              f"Cmp is derived as ROUND(Att*Cmp_pct/100,0); Group is derived from Squad.").font = NOTE_FONT

# ---------------------------------------------------------------------
# Sheet: Sample (n=60)
# ---------------------------------------------------------------------
ws = wb.create_sheet("Sample (n=60)")
for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
style_header_row(ws, 1, len(headers))
ws.freeze_panes = "A2"

for i, row in sample.iterrows():
    r = i + 2
    ws.cell(row=r, column=1, value=row["Player"]).font = BODY_FONT
    ws.cell(row=r, column=2, value=row["Squad"]).font = BODY_FONT
    ws.cell(row=r, column=3, value=row["Pos"]).font = BODY_FONT
    ws.cell(row=r, column=4, value=int(row["Att"])).font = BODY_FONT
    ws.cell(row=r, column=5, value=int(row["Cmp_pct"])).font = BODY_FONT
    ws.cell(row=r, column=6, value=f"=ROUND(D{r}*E{r}/100,0)").font = BODY_FONT
    ws.cell(row=r, column=7, value=f'=IF(OR(B{r}="ESP",B{r}="ARG"),"Finalist","Non-finalist")').font = BODY_FONT
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = BORDER
    if row["Group"] == "Finalist":
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="EAF1FB")
    else:
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FDF1EC")

autosize(ws, [22, 8, 6, 8, 10, 14, 14])
n_fin = len(fin_sample)
n_non = len(non_sample)
FIN_FIRST, FIN_LAST = 2, 1 + n_fin
NON_FIRST, NON_LAST = 2 + n_fin, 1 + n_fin + n_non

note_row = NON_LAST + 2
ws.cell(row=note_row, column=1,
        value=f"Simple random sample drawn in Python (pandas .sample, random_state={RANDOM_SEED}, "
              f"the student ID, for reproducibility) from the wrangled population on the "
              f"'Data (n=450)' sheet. Rows {FIN_FIRST}-{FIN_LAST} = Finalist sample (n={n_fin}); "
              f"rows {NON_FIRST}-{NON_LAST} = Non-finalist sample (n={n_non}).").font = NOTE_FONT

SAMPLE_SHEET = "'Sample (n=60)'"
FIN_RANGE = f"{SAMPLE_SHEET}!$E${FIN_FIRST}:$E${FIN_LAST}"
NON_RANGE = f"{SAMPLE_SHEET}!$E${NON_FIRST}:$E${NON_LAST}"

# ---------------------------------------------------------------------
# Sheet: Descriptive Stats
# ---------------------------------------------------------------------
ws = wb.create_sheet("Descriptive Stats")
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Descriptive statistics — Cmp% (pass completion %)").font = TITLE_FONT
ws.cell(row=2, column=1, value="All formulas reference the 'Sample (n=60)' sheet directly.").font = SUBTITLE_FONT

headers2 = ["Group", "n", "Mean", "Median", "Std Dev (sample)", "Min", "Max"]
hdr_row = 4
for c, h in enumerate(headers2, start=1):
    ws.cell(row=hdr_row, column=c, value=h)
style_header_row(ws, hdr_row, len(headers2))

rows2 = [
    ("Finalist (Spain + Argentina)", FIN_RANGE),
    ("Non-finalist", NON_RANGE),
]
for i, (label, rng) in enumerate(rows2):
    r = hdr_row + 1 + i
    ws.cell(row=r, column=1, value=label).font = BOLD_FONT
    ws.cell(row=r, column=2, value=f"=COUNT({rng})").font = BODY_FONT
    ws.cell(row=r, column=3, value=f"=AVERAGE({rng})").font = BODY_FONT
    ws.cell(row=r, column=3).number_format = "0.00"
    ws.cell(row=r, column=4, value=f"=MEDIAN({rng})").font = BODY_FONT
    ws.cell(row=r, column=5, value=f"=STDEV({rng})").font = BODY_FONT
    ws.cell(row=r, column=5).number_format = "0.00"
    ws.cell(row=r, column=6, value=f"=MIN({rng})").font = BODY_FONT
    ws.cell(row=r, column=7, value=f"=MAX({rng})").font = BODY_FONT
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = BORDER

autosize(ws, [30, 8, 10, 10, 16, 8, 8])
DESC_FIN_ROW = hdr_row + 1
DESC_NON_ROW = hdr_row + 2

# ---------------------------------------------------------------------
# Sheet: Inferential Stats
# ---------------------------------------------------------------------
ws = wb.create_sheet("Inferential Stats")
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Inferential statistics — Finalist vs Non-finalist Cmp%").font = TITLE_FONT
ws.cell(row=2, column=1,
        value="95% CI for the difference in means (Welch-Satterthwaite) and Welch's two-sample t-test.").font = SUBTITLE_FONT

DS = "'Descriptive Stats'"
labels_vals = [
    ("n (Finalist)",        f"={DS}!B{DESC_FIN_ROW}",  None),
    ("n (Non-finalist)",    f"={DS}!B{DESC_NON_ROW}",  None),
    ("Mean (Finalist)",     f"={DS}!C{DESC_FIN_ROW}",  "0.00"),
    ("Mean (Non-finalist)", f"={DS}!C{DESC_NON_ROW}",  "0.00"),
    ("Variance (Finalist), sample",     f"=VAR({FIN_RANGE})", "0.00"),
    ("Variance (Non-finalist), sample", f"=VAR({NON_RANGE})", "0.00"),
]
r = 4
rowref = {}
for label, formula, fmt in labels_vals:
    ws.cell(row=r, column=1, value=label).font = BODY_FONT
    cell = ws.cell(row=r, column=2, value=formula)
    cell.font = BODY_FONT
    if fmt:
        cell.number_format = fmt
    rowref[label] = r
    r += 1

r += 1
n1 = f"B{rowref['n (Finalist)']}"
n2 = f"B{rowref['n (Non-finalist)']}"
m1 = f"B{rowref['Mean (Finalist)']}"
m2 = f"B{rowref['Mean (Non-finalist)']}"
v1 = f"B{rowref['Variance (Finalist), sample']}"
v2 = f"B{rowref['Variance (Non-finalist), sample']}"

calc_rows = []
calc_rows.append(("Mean difference (Finalist − Non-finalist)", f"={m1}-{m2}", "0.00"))
diff_row = r
r_diff = r
r += 1
calc_rows.append(("Standard error of difference", f"=SQRT({v1}/{n1}+{v2}/{n2})", "0.00"))
se_row = r
r += 1
calc_rows.append(
    ("Welch-Satterthwaite df",
     f"=(({v1}/{n1}+{v2}/{n2})^2)/((({v1}/{n1})^2)/({n1}-1)+(({v2}/{n2})^2)/({n2}-1))",
     "0.0")
)
df_row = r
r += 1
calc_rows.append(("t critical (two-tailed, alpha=0.05)", f"=TINV(0.05,B{df_row})", "0.000"))
tcrit_row = r
r += 1
calc_rows.append(("Margin of error", f"=B{tcrit_row}*B{se_row}", "0.00"))
margin_row = r
r += 1
calc_rows.append(("95% CI lower bound", f"=B{r_diff}-B{margin_row}", "0.00"))
lo_row = r
r += 1
calc_rows.append(("95% CI upper bound", f"=B{r_diff}+B{margin_row}", "0.00"))
hi_row = r
r += 1
calc_rows.append(("t-statistic", f"=B{r_diff}/B{se_row}", "0.000"))
tstat_row = r
r += 1
calc_rows.append(
    ("p-value (Welch's two-sample t-test, two-tailed)",
     f"=TTEST({FIN_RANGE},{NON_RANGE},2,3)", "0.0000")
)
pval_row = r
r += 1
calc_rows.append(("alpha", "=0.05", "0.00"))
alpha_row = r
r += 1

start_calc = diff_row
for i, (label, formula, fmt) in enumerate(calc_rows):
    rr = start_calc + i
    ws.cell(row=rr, column=1, value=label).font = BOLD_FONT if i in (0, 8) else BODY_FONT
    cell = ws.cell(row=rr, column=2, value=formula)
    cell.font = BOLD_FONT if i in (0, 8) else BODY_FONT
    if fmt:
        cell.number_format = fmt

r += 1
ws.cell(row=r, column=1, value="Decision at alpha = 0.05").font = BOLD_FONT
ws.cell(row=r, column=2,
        value=f'=IF(B{pval_row}<B{alpha_row},"Reject H0 - significant difference",'
              f'"Fail to reject H0 - not statistically significant")').font = BOLD_FONT
decision_row = r
r += 2
ws.cell(row=r, column=1,
        value="H0: mean Cmp%(Finalist) = mean Cmp%(Non-finalist)  |  "
              "H1: mean Cmp%(Finalist) != mean Cmp%(Non-finalist)").font = NOTE_FONT

autosize(ws, [42, 16])
for rr in list(range(4, r + 1)):
    ws.cell(row=rr, column=1).border = BORDER
    ws.cell(row=rr, column=2).border = BORDER

# ---------------------------------------------------------------------
# Sheet: Answer Chart (embedded PNG for reference, not a live formula)
# ---------------------------------------------------------------------
ws = wb.create_sheet("Answer Chart")
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Reference chart (static image, generated in Python/matplotlib)").font = SUBTITLE_FONT
try:
    img = XLImage("task3_answer_card.png")
    img.width = img.width * 0.55
    img.height = img.height * 0.55
    ws.add_image(img, "A3")
except Exception as e:
    ws.cell(row=3, column=1, value=f"(chart image not embedded: {e})").font = NOTE_FONT

wb.save("task3_passing.xlsx")
print("saved task3_passing.xlsx")
print("FIN_RANGE", FIN_RANGE, "NON_RANGE", NON_RANGE)
